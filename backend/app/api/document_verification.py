"""Prototype document evidence screening. This is not issuer authentication."""
import io
import re
from pathlib import Path
from typing import Any

from docx import Document as DocxDocument
from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from openpyxl import load_workbook

from app.core.security import require_user
from app.api.tender_document_extraction import _extract_text

router = APIRouter(prefix="/api/v1/document-verification", tags=["Prototype document verification"])
MAX_BYTES = 15 * 1024 * 1024
ALLOWED = {".pdf", ".png", ".jpg", ".jpeg", ".docx", ".xlsx", ".txt"}
DOCUMENT_TYPES = {
    "company_registration", "company_profile", "gst_certificate", "pan_card",
    "turnover_certificate", "solvency_certificate", "work_completion_certificate",
    "electrical_contractor_licence", "manufacturing_facility", "tender_fee_proof",
    "emd_proof", "other",
}
LABELS = {
    "company_registration": "Company registration details",
    "company_profile": "Company profile",
    "gst_certificate": "GST registration certificate",
    "pan_card": "PAN card",
    "turnover_certificate": "Turnover certificate (Annexure 1)",
    "solvency_certificate": "Solvency certificate",
    "work_completion_certificate": "Similar-work completion certificate",
    "electrical_contractor_licence": "A-Class electrical contractor licence",
    "manufacturing_facility": "Manufacturing facility details",
    "tender_fee_proof": "Tender fee evidence",
    "emd_proof": "EMD evidence",
    "other": "Other supporting document",
}


def _document_text(content: bytes, extension: str) -> tuple[str, str]:
    if extension in {".pdf", ".png", ".jpg", ".jpeg"}:
        return _extract_text(content, extension)
    try:
        if extension == ".txt":
            text = content.decode("utf-8-sig", errors="replace")[:250_000]
        elif extension == ".docx":
            document = DocxDocument(io.BytesIO(content))
            parts = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            text = "\n".join(parts).strip()
        elif extension == ".xlsx":
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            parts = []
            for sheet in workbook.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    values = [str(value).strip() for value in row if value is not None and str(value).strip()]
                    if values:
                        parts.append(" | ".join(values))
                    if len(parts) >= 5000:
                        break
            workbook.close()
            text = "\n".join(parts).strip()
        else:
            raise HTTPException(status_code=415, detail="Unsupported file type.")
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=422, detail="The file could not be parsed. Try exporting it again.") from exc
    if not text:
        raise HTTPException(status_code=422, detail="No readable text was found in this document.")
    return text[:250_000], "document_parser"


def _first(pattern: str, text: str, flags: int = re.IGNORECASE) -> str | None:
    match = re.search(pattern, text, flags)
    return match.group(1).strip(" \\t:;,.") if match else None


def _extract_fields(text: str) -> dict[str, Any]:
    pan = _first(r"\b([A-Z]{5}[0-9]{4}[A-Z])\b", text, re.IGNORECASE)
    gstin = _first(r"\b([0-9]{2}[A-Z]{5}[0-9]{4}[A-Z][1-9A-Z]Z[0-9A-Z])\b", text, re.IGNORECASE)
    registration_no = _first(
        r"(?:CIN|LLPIN|registration\s*(?:no\.?|number)|company\s*no\.?)\s*[:#\-]?\s*([A-Z0-9/\-]{5,30})",
        text,
    )
    business_name = _first(
        r"(?:legal\s+name|name\s+of\s+(?:the\s+)?(?:company|business|firm|proprietor|taxpayer)|company\s+name|business\s+name)\s*[:\-]\s*([^\n\r]{2,160})",
        text,
    )
    money_values = []
    for match in re.finditer(r"(?:₹|INR\s*|Rs\.?\s*)([0-9][0-9,]*(?:\.[0-9]+)?)\s*(crore|crores|lakh|lakhs|lac|lacs|thousand)?", text, re.IGNORECASE):
        try:
            number = float(match.group(1).replace(",", ""))
        except ValueError:
            continue
        unit = (match.group(2) or "").lower()
        multiplier = 10_000_000 if unit.startswith("crore") else 100_000 if unit.startswith(("lakh", "lac")) else 1_000 if unit == "thousand" else 1
        money_values.append({"display": match.group(0).strip(), "value_inr": int(number * multiplier)})
        if len(money_values) >= 20:
            break
    expiry = _first(r"(?:valid\s*(?:until|upto|through)|expiry\s*date|validity\s*date|expires\s*on)\s*[:\-]?\s*([0-3]?\d[\/\-.][01]?\d[\/\-.](?:\d{2}|\d{4})|\d{4}[\/\-.][01]?\d[\/\-.][0-3]?\d)", text)
    return {
        "business_name": business_name[:160] if business_name else None,
        "pan": pan.upper() if pan else None,
        "gstin": gstin.upper() if gstin else None,
        "registration_number": registration_no.upper() if registration_no else None,
        "expiry_date_text": expiry,
        "amounts": money_values,
    }


def _checks(document_type: str, text: str, fields: dict[str, Any]) -> list[dict[str, str]]:
    checks: list[dict[str, str]] = [{
        "code": "readable_text",
        "label": "Readable document text",
        "status": "pass" if len(text.strip()) >= 30 else "manual_review",
        "detail": "Text was extracted for review." if len(text.strip()) >= 30 else "Very little text was extracted; inspect the original file.",
    }]
    if document_type == "pan_card":
        present = bool(fields["pan"])
        checks.append({"code": "pan_format", "label": "PAN format", "status": "pass" if present else "manual_review", "detail": "A PAN-shaped identifier was extracted; issuer status is not checked." if present else "PAN could not be reliably extracted."})
    elif document_type == "gst_certificate":
        present = bool(fields["gstin"])
        checks.append({"code": "gstin_format", "label": "GSTIN format", "status": "pass" if present else "manual_review", "detail": "A GSTIN-shaped identifier was extracted; GST registration status is not checked." if present else "GSTIN could not be reliably extracted."})
    elif document_type == "company_registration":
        present = bool(fields["registration_number"])
        checks.append({"code": "registration_number", "label": "Registration number", "status": "pass" if present else "manual_review", "detail": "A registration identifier was extracted; the registry was not queried." if present else "Registration number could not be reliably extracted."})
    elif document_type in {"turnover_certificate", "solvency_certificate", "work_completion_certificate", "emd_proof", "tender_fee_proof"}:
        present = bool(fields["amounts"])
        checks.append({"code": "numeric_amount", "label": "Numeric amount", "status": "pass" if present else "manual_review", "detail": "One or more explicit numeric amounts were extracted; confirm the correct amount and context." if present else "No unambiguous numeric amount was extracted."})
    else:
        checks.append({"code": "required_fields", "label": "Required fields", "status": "manual_review", "detail": "This document type needs reviewer-defined checks."})
    checks.append({
        "code": "business_name_extracted", "label": "Business name",
        "status": "pass" if fields.get("business_name") else "manual_review",
        "detail": "A possible business/legal name was extracted. Compare it with authoritative records." if fields.get("business_name") else "Business name was not confidently identified by the prototype parser.",
    })
    checks.append({"code": "issuer_authenticity", "label": "Issuer authenticity", "status": "not_checked", "detail": "No government registry, issuer API, DigiLocker or EntityLocker connection was used."})
    checks.append({"code": "signature_and_tampering", "label": "Digital signature / tampering", "status": "not_checked", "detail": "This prototype does not validate digital signatures or prove that a document has not been altered."})
    return checks


@router.get("/document-types")
def document_types(user: dict = Depends(require_user)) -> dict[str, Any]:
    return {"items": [{"value": key, "label": value} for key, value in LABELS.items()], "mode": "prototype"}


@router.post("/analyze")
async def analyze_document(
    document_type: str = Form(...),
    file: UploadFile = File(...),
    user: dict = Depends(require_user),
) -> dict[str, Any]:
    """Analyze an uploaded file transiently; the original file is not persisted by this endpoint."""
    if document_type not in DOCUMENT_TYPES:
        raise HTTPException(status_code=422, detail="Choose a supported document type.")
    filename = Path(file.filename or "document").name
    extension = Path(filename).suffix.lower()
    if extension not in ALLOWED:
        raise HTTPException(status_code=415, detail="Allowed types: PDF, PNG, JPG, DOCX, XLSX, TXT.")
    content = await file.read(MAX_BYTES + 1)
    await file.close()
    if not content:
        raise HTTPException(status_code=400, detail="The selected file is empty.")
    if len(content) > MAX_BYTES:
        raise HTTPException(status_code=413, detail="Maximum file size is 15 MiB.")
    text, method = _document_text(content, extension)
    fields = _extract_fields(text)
    checks = _checks(document_type, text, fields)
    manual_review = any(item["status"] == "manual_review" for item in checks)
    return {
        "mode": "prototype",
        "filename": filename,
        "document_type": document_type,
        "document_label": LABELS[document_type],
        "extraction_method": method,
        "text_characters": len(text),
        "fields": fields,
        "checks": checks,
        "overall_status": "manual_review" if manual_review else "extracted_unverified",
        "issuer_verified": False,
        "stored": False,
        "review_required": True,
        "disclaimer": "Prototype evidence screening only. Passing format or extraction checks does not establish authenticity, issuer status, legal validity, or tender eligibility. No DigiLocker/EntityLocker or government registry was queried. Confirm the original document and authoritative source manually.",
    }
