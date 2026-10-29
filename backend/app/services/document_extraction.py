"""Private document content retrieval and bounded text extraction for screening."""
import io
from urllib.request import Request, urlopen

from fastapi import HTTPException
from docx import Document as WordDocument
from openpyxl import load_workbook

from app.api.tender_document_extraction import _extract_text
from app.core.cloudinary_storage import get_signed_download_url

MAX_DOCUMENT_BYTES = 15 * 1024 * 1024
MAX_TEXT_CHARS = 250_000


def _download_content(metadata: dict) -> bytes:
    public_id = metadata.get("cloudinaryPublicId")
    if not public_id:
        raise ValueError("Document is not stored in the current Cloudinary storage.")
    signed_url = get_signed_download_url(str(public_id))
    request = Request(signed_url, headers={"User-Agent": "NIRNAY/1.0"})
    with urlopen(request, timeout=20) as response:
        content = response.read(MAX_DOCUMENT_BYTES + 1)
    if len(content) > MAX_DOCUMENT_BYTES:
        raise ValueError("Document exceeds the 15 MiB processing limit.")
    if not content:
        raise ValueError("Document is empty.")
    return content


def extract_document_text(metadata: dict) -> str:
    """Retrieve a user's private file and extract searchable text without persisting it."""
    filename = str(metadata.get("filename") or "document")
    extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    content = _download_content(metadata)

    if extension in {"pdf", "png", "jpg", "jpeg"}:
        text, _method = _extract_text(content, "." + extension)
    elif extension == "txt":
        text = content.decode("utf-8-sig", errors="replace")
    elif extension == "docx":
        document = WordDocument(io.BytesIO(content))
        parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        text = "\n".join(parts)
    elif extension == "xlsx":
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        parts = []
        try:
            for sheet in workbook.worksheets:
                parts.append(f"Sheet: {sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    values = [str(value).strip() for value in row if value is not None and str(value).strip()]
                    if values:
                        parts.append(" | ".join(values))
                    if len("\n".join(parts)) >= MAX_TEXT_CHARS:
                        break
                if len("\n".join(parts)) >= MAX_TEXT_CHARS:
                    break
        finally:
            workbook.close()
        text = "\n".join(parts)
    else:
        raise ValueError("Unsupported file type for content screening.")

    text = text.strip()
    if not text:
        raise ValueError("No readable text could be extracted from this document.")
    return text[:MAX_TEXT_CHARS]
