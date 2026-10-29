import io

from docx import Document
from openpyxl import Workbook

from app.services import document_extraction


def test_extract_text_file(monkeypatch):
    monkeypatch.setattr(document_extraction, "_download_content", lambda metadata: b"GST certificate valid through 2027")
    result = document_extraction.extract_document_text({"filename": "gst.txt"})
    assert "GST certificate" in result


def test_extract_docx_paragraphs_and_tables(monkeypatch):
    stream = io.BytesIO()
    document = Document()
    document.add_paragraph("GST registration certificate")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Turnover"
    table.cell(0, 1).text = "₹50 lakh"
    document.save(stream)
    monkeypatch.setattr(document_extraction, "_download_content", lambda metadata: stream.getvalue())
    result = document_extraction.extract_document_text({"filename": "evidence.docx"})
    assert "GST registration certificate" in result
    assert "Turnover | ₹50 lakh" in result


def test_extract_xlsx_cells(monkeypatch):
    stream = io.BytesIO()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Financials"
    sheet.append(["Metric", "Value"])
    sheet.append(["Annual turnover", "5000000"])
    workbook.save(stream)
    monkeypatch.setattr(document_extraction, "_download_content", lambda metadata: stream.getvalue())
    result = document_extraction.extract_document_text({"filename": "financials.xlsx"})
    assert "Sheet: Financials" in result
    assert "Annual turnover | 5000000" in result


def test_empty_text_document_is_rejected(monkeypatch):
    monkeypatch.setattr(document_extraction, "_download_content", lambda metadata: b"  ")
    try:
        document_extraction.extract_document_text({"filename": "empty.txt"})
    except ValueError as exc:
        assert "No readable text" in str(exc)
    else:
        raise AssertionError("Empty text should not be treated as evidence")
